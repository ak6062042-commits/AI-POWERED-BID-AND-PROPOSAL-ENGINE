import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def build_proposal_docx(workspace: dict, output_path: str) -> str:
    doc = Document()

    doc.add_heading(workspace.get("name", "Proposal Response"), level=0)
    doc.add_paragraph(f"Workspace ID: {workspace['id']}")
    doc.add_paragraph(f"Generated: {workspace['created_at']}")
    doc.add_paragraph(
        "This document was auto-drafted by the AI Bid & Proposal Response "
        "Engine. All sections marked [needs_review] should be checked by "
        "the bid manager before final submission."
    )

    go_no_go = workspace.get("go_no_go") or {}
    scoring = workspace.get("scoring_result") or {}
    if go_no_go or scoring:
        _add_heading(doc, "Bid Opportunity Assessment", level=1)
        p = doc.add_paragraph()
        run = p.add_run(f"Recommendation: {go_no_go.get('decision', 'N/A')}")
        run.bold = True
        run.font.size = Pt(13)
        if scoring:
            doc.add_paragraph(
                f"Modeled win probability: {scoring.get('win_probability_pct')}% "
                f"(historical base rate: {scoring.get('base_win_rate_pct')}%, "
                f"model accuracy on training data: {scoring.get('model_train_accuracy_pct')}%)"
            )
        for reason in go_no_go.get("reasons", []):
            doc.add_paragraph(f"• {reason}", style="List Bullet")

    checklist = workspace.get("checklist") or []
    summary = workspace.get("checklist_summary") or {}
    if checklist:
        _add_heading(doc, "Compliance Checklist", level=1)
        doc.add_paragraph(
            f"{summary.get('pass', 0)} pass / {summary.get('partial', 0)} partial / "
            f"{summary.get('gap', 0)} gap out of {summary.get('total_requirements', 0)} "
            f"requirements ({summary.get('requirements_matched_pct', 0)}% matched)."
        )
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Requirement"
        hdr[1].text = "Status"
        hdr[2].text = "Best supporting evidence"
        for item in checklist:
            row = table.add_row().cells
            row[0].text = item["requirement"][:300]
            row[1].text = item["status"].upper()
            best = item.get("best_match")
            row[2].text = best["title"] if best else "No matching evidence found"

    draft = workspace.get("draft") or []
    if draft:
        _add_heading(doc, "Proposal Response", level=1)
        for section in draft:
            doc.add_heading(f"{section.get('id')}: {section.get('question')}", level=2)
            doc.add_paragraph(section.get("draft_response", ""))
            evidence = section.get("evidence_used") or []
            if evidence:
                ev_text = ", ".join(e["title"] for e in evidence)
                p = doc.add_paragraph(f"Evidence referenced: {ev_text}")
                for run in p.runs:
                    run.italic = True
                    run.font.size = Pt(9)

    doc.save(output_path)
    return output_path
