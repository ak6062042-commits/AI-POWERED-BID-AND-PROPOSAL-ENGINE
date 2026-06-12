
import os
import re
import json
from dataclasses import dataclass, field, asdict
from typing import List, Dict, Any

import pdfplumber
from docx import Document as DocxDocument

@dataclass
class ExtractionResult:
    full_text: str
    deadlines: List[Dict[str, str]] = field(default_factory=list)
    budget: List[Dict[str, str]] = field(default_factory=list)
    mandatory_requirements: List[Dict[str, Any]] = field(default_factory=list)
    evaluation_criteria: List[Dict[str, Any]] = field(default_factory=list)
    qa_sections: List[Dict[str, str]] = field(default_factory=list)
    enrichment_used: bool = False

    def to_dict(self) -> Dict[str, Any]:
        d = asdict(self)
        d.pop("full_text", None)
        return d

def load_text(file_path: str) -> str:
    lower = file_path.lower()
    if lower.endswith(".pdf"):
        text_parts = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                text_parts.append(page_text)
        return "\n".join(text_parts)
    elif lower.endswith(".docx"):
        doc = DocxDocument(file_path)
        parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                parts.append(" | ".join(cells))
        return "\n".join(parts)
    elif lower.endswith(".txt"):
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    else:
        raise ValueError(f"Unsupported file type: {file_path}")


DATE_PATTERN = re.compile(
    r"\b(\d{1,2}(?:st|nd|rd|th)?\s+(?:January|February|March|April|May|June|July|"
    r"August|September|October|November|December)\s+\d{4}|"
    r"\d{1,2}[/-]\d{1,2}[/-]\d{2,4})\b",
    re.IGNORECASE,
)

DEADLINE_KEYWORDS = [
    "deadline", "submission", "submit", "no later than", "closing date",
    "due date", "last date", "must be submitted",
]

BUDGET_KEYWORDS = [
    "budget", "estimated cost", "contract value", "ceiling", "PKR", "Rs.",
    "PKR.", "USD", "financial proposal ceiling",
]

BUDGET_AMOUNT_PATTERN = re.compile(
    r"(PKR|Rs\.?|USD)\s?[\d,]+(?:\.\d+)?(?:\s?(?:million|billion|m|bn))?",
    re.IGNORECASE,
)

MANDATORY_KEYWORDS = re.compile(
    r"\b(shall|must|mandatory|required to|is required|are required|"
    r"will be required)\b",
    re.IGNORECASE,
)

WEIGHT_PATTERN = re.compile(r"(.+?)[\-–—:]\s*(\d{1,3})\s?%")

QUESTION_PATTERN = re.compile(
    r"^(Q\d+|Question\s*\d+)\s*[:\.\-]\s*(.+)", re.IGNORECASE
)


HEADING_PATTERN = re.compile(r"^\d{1,2}\.\s+[A-Z][\w\s,/&\-]{2,60}$")


def _is_heading(line: str) -> bool:
    """Detect numbered section headings like '4. Mandatory Eligibility...'"""
    stripped = line.strip()
    if HEADING_PATTERN.match(stripped) and not stripped.rstrip().endswith((".", ":")):
        return True
    if stripped.isupper() and len(stripped.split()) <= 8:
        return True
    return False


def _split_sentences(text: str) -> List[str]:
    lines = []
    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            continue
        if len(line) > 220:
            lines.extend(re.split(r"(?<=[.;])\s+", line))
        else:
            lines.append(line)
    return [l.strip() for l in lines if l.strip()]


def extract_deadlines(lines: List[str]) -> List[Dict[str, str]]:
    results = []
    for line in lines:
        if _is_heading(line):
            continue
        lower = line.lower()
        if any(k in lower for k in DEADLINE_KEYWORDS):
            dates = DATE_PATTERN.findall(line)
            if dates:
                for d in dates:
                    results.append({"date": d.strip(), "context": line[:300]})
            else:
                results.append({"date": "", "context": line[:300]})
    return results


def extract_budget(lines: List[str]) -> List[Dict[str, str]]:
    results = []
    for line in lines:
        if _is_heading(line):
            continue
        lower = line.lower()
        if any(k.lower() in lower for k in BUDGET_KEYWORDS):
            amounts = BUDGET_AMOUNT_PATTERN.findall(line)
            full_matches = BUDGET_AMOUNT_PATTERN.finditer(line)
            amount_strs = [m.group(0) for m in full_matches]
            if amount_strs:
                for a in amount_strs:
                    results.append({"amount": a.strip(), "context": line[:300]})
            elif "budget" in lower or "ceiling" in lower:
                results.append({"amount": "", "context": line[:300]})
    return results


def extract_mandatory_requirements(lines: List[str]) -> List[Dict[str, Any]]:
    results = []
    req_id = 1
    for line in lines:

        if len(line) < 15 or _is_heading(line):
            continue
        if MANDATORY_KEYWORDS.search(line):

            if WEIGHT_PATTERN.match(line):
                continue
            results.append({
                "id": f"REQ-{req_id:03d}",
                "text": line,
            })
            req_id += 1
    return results


def extract_evaluation_criteria(lines: List[str]) -> List[Dict[str, Any]]:
    results = []
    for line in lines:
        m = WEIGHT_PATTERN.match(line)
        if m:
            criterion = m.group(1).strip().lstrip("•-*").strip()
            weight = int(m.group(2))
      
            if 0 < weight <= 100 and len(criterion) > 3:
                results.append({"criterion": criterion, "weight_pct": weight})
    return results


def extract_qa_sections(lines: List[str]) -> List[Dict[str, str]]:
    results = []
    for line in lines:
        m = QUESTION_PATTERN.match(line)
        if m:
            results.append({"id": m.group(1).upper().replace(" ", ""),
                             "question": m.group(2).strip()})
        elif line.strip().endswith("?") and len(line) > 20:
            results.append({"id": f"Q-AUTO-{len(results)+1}",
                             "question": line.strip()})
    return results


def _llm_enrich(result: ExtractionResult) -> ExtractionResult:
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        return result

    try:
        import anthropic
        client = anthropic.Anthropic(api_key=api_key)

        excerpt = result.full_text[:12000]

        prompt = f"""You are assisting a bid manager. Below is raw text extracted
from an RFP/Tender document. I have already run a rule-based extractor and
got the JSON below. Your job:

1. Review the "mandatory_requirements" list and remove any that are not
   genuinely mandatory compliance clauses (e.g. remove duplicates or
   boilerplate). Keep IDs as given.
2. For each remaining requirement, add a "short_label" field: a 3-6 word
   plain-English label summarising the requirement (e.g. "ISO 27001
   certification required").
3. Review "evaluation_criteria" and fix any obviously wrong weight numbers.
4. Return ONLY valid JSON, with the exact same top-level keys as the input
   JSON (mandatory_requirements, evaluation_criteria), nothing else --
   no markdown fences, no commentary.

RFP TEXT EXCERPT:
---
{excerpt}
---

CURRENT JSON:
{json.dumps({"mandatory_requirements": result.mandatory_requirements, "evaluation_criteria": result.evaluation_criteria})}
"""

        response = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=4000,
            messages=[{"role": "user", "content": prompt}],
        )
        text_out = "".join(
            block.text for block in response.content if hasattr(block, "text")
        ).strip()
        text_out = re.sub(r"^```(?:json)?|```$", "", text_out, flags=re.MULTILINE).strip()
        enriched = json.loads(text_out)

        if "mandatory_requirements" in enriched:
            result.mandatory_requirements = enriched["mandatory_requirements"]
        if "evaluation_criteria" in enriched and enriched["evaluation_criteria"]:
            result.evaluation_criteria = enriched["evaluation_criteria"]
        result.enrichment_used = True
    except Exception as e: 
        result.enrichment_used = False
        result.deadlines.append({"date": "", "context": f"[LLM enrichment skipped: {e}]"})

    return result


def extract(file_path: str, use_llm: bool = True) -> ExtractionResult:
    text = load_text(file_path)
    lines = _split_sentences(text)

    result = ExtractionResult(
        full_text=text,
        deadlines=extract_deadlines(lines),
        budget=extract_budget(lines),
        mandatory_requirements=extract_mandatory_requirements(lines),
        evaluation_criteria=extract_evaluation_criteria(lines),
        qa_sections=extract_qa_sections(lines),
    )

    if use_llm:
        result = _llm_enrich(result)

    return result


if __name__ == "__main__":
    import sys
    path = sys.argv[1] if len(sys.argv) > 1 else \
        "/home/claude/bidengine/sample_rfps/PITB_ICMS_RFP_2026.docx"
    r = extract(path, use_llm=False)
    print(json.dumps(r.to_dict(), indent=2))
